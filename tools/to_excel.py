# -*- coding: utf-8 -*-
"""把 docs/data.json 导成 WPS / Excel 能直接打开的表格。

用法： python tools/to_excel.py
产物： 纳指溢价_YYYYMMDD_HHMM.xlsx
"""
from __future__ import annotations

import json
import os
import sys
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("缺少 openpyxl：pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def main() -> None:
    path = os.path.join(ROOT, "docs", "data.json")
    if not os.path.exists(path):
        sys.exit("还没有 docs/data.json，先运行 python run.py")
    with open(path, encoding="utf-8") as fh:
        data = json.load(fh)

    funds = [f for f in data["funds"] if f.get("ok")]
    wb = Workbook()

    ws = wb.active
    ws.title = "实时溢价"
    cols = ["代码", "名称", "类型", "现价", "今日涨跌", "估算净值", "最新净值",
            "净值日期", "溢价率", "20日均溢价", "距均值", "60日分位",
            "年内涨幅", "近20日超额", "成交额(亿)", "信号"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center")

    for f in funds:
        ws.append([
            f["code"], f["name"], f["type"], f["price"], f["pct"], f["est_nav"], f["nav"],
            f["nav_date"], f["premium"], f["mean20"], f["spread20"], f["pct60"],
            f["ret_ytd"], f["rs_20d"], f["amount"] / 1e8, f["signal"],
        ])

    pct_cols = {4, 5, 6, 7, 9, 10, 11, 13, 14}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for idx in pct_cols:
            row[idx - 1].number_format = "0.00%"
        row[11].number_format = "0%"
        row[14].number_format = "0.00"
    for i, w in enumerate([8, 20, 6, 8, 9, 10, 10, 11, 9, 11, 9, 9, 9, 11, 11, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # 溢价历史矩阵
    ws2 = wb.create_sheet("溢价历史")
    dates: list[str] = []
    series: dict[str, dict[str, float]] = {}
    for f in funds:
        series[f["name"]] = {h["d"]: h["premium"] for h in f.get("history", [])}
        dates.extend(series[f["name"]].keys())
    dates = sorted(set(dates))
    ws2.append(["日期"] + list(series.keys()))
    for c in range(1, len(series) + 2):
        cell = ws2.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for d in dates:
        ws2.append([d] + [series[n].get(d) for n in series])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=2):
        for cell in row:
            cell.number_format = "0.0%"
    ws2.column_dimensions["A"].width = 12
    ws2.freeze_panes = "B2"

    out = os.path.join(ROOT, f"纳指溢价_{datetime.now():%Y%m%d_%H%M}.xlsx")
    wb.save(out)
    print(f"已生成 {out}")


if __name__ == "__main__":
    main()
